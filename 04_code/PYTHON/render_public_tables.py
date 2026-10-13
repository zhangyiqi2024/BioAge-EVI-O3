"""Rebuild all manuscript tables from public aggregate CSV and JSON sources.

Run from any working directory. This renderer needs only 03_source_data/tables
and the Python standard library. No DOCX, participant data or model fit is used.
An optional --xlsx workbook requires openpyxl. Displayed numerical strings are
preserved, and the generated HTML is parsed back for a cell-level comparison.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import html
from html.parser import HTMLParser
import json
import re
import shutil
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
CSS = """
:root { color-scheme: light; }
body { margin: 2rem auto; padding: 0 1.25rem; max-width: 100rem;
  color: #18222d; background: #fff; font: 15px/1.5 Arial, Helvetica, sans-serif; }
h1,h2,h3 { line-height: 1.25; }
h1 { font-size: 1.7rem; } h2 { font-size: 1.3rem; }
a { color: #185991; } .scope { border-left: 4px solid #366482; padding: .6rem 1rem; background: #f2f6f9; }
.table-wrap { overflow-x: auto; margin: 1rem 0; }
table { border-collapse: collapse; width: 100%; font-size: 13px; line-height: 1.35; }
caption { text-align: left; font-weight: 700; padding: 0 0 .7rem; font-size: 15px; }
thead { border-top: 2px solid #263745; border-bottom: 1px solid #263745; }
th,td { padding: .5rem .6rem; vertical-align: top; border-bottom: 1px solid #d7dfe4;
  overflow-wrap: anywhere; white-space: pre-wrap; }
thead th { background: #eef3f6; text-align: left; }
tbody th { text-align: left; font-weight: 500; }
td.numeric { text-align: right; font-variant-numeric: tabular-nums; }
tr.group-row th { background: #eef3f6; font-weight: 700; }
tbody tr:last-child th,tbody tr:last-child td { border-bottom: 2px solid #263745; }
.notes p { margin: .55rem 0; font-size: 13px; }
.provenance { margin: 1rem 0; padding-top: .5rem; border-top: 1px solid #d7dfe4; font-size: 12px; color: #435466; }
.table-section { margin: 2rem 0 3rem; }
.downloads { display: flex; gap: 1rem; margin: .5rem 0; }
input[type=search] { width: min(38rem,90%); padding: .6rem; font: inherit; }
li { margin: .35rem 0; } .muted { color: #536571; }
@media print {
  @page { size: A4 landscape; margin: 12mm; }
  body { margin: 0; padding: 0; max-width: none; color: black; font-size: 10pt; }
  .navigation,.downloads,.search { display: none; }
  .table-section { break-before: page; margin: 0; }
  .table-section:first-of-type { break-before: auto; }
  .table-wrap { overflow: visible; }
  table { font-size: 8pt; } caption { font-size: 10pt; }
  th,td { padding: 3pt 4pt; }
  thead { display: table-header-group; }
  tr { break-inside: avoid; } h2,h3,caption { break-after: avoid; }
  .notes p,.provenance { font-size: 8pt; }
}
"""
SCOPE = (
    "These tables reconstruct the authors' archived displayed results from public "
    "aggregate sources. They do not refit models or independently validate estimates. "
    "No participant records are required."
)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def safe_child(directory: Path, name: str) -> Path:
    path = (directory / name).resolve()
    try:
        path.relative_to(directory.resolve())
    except ValueError as exc:
        raise ValueError(f"Source path escapes the table directory: {name}") from exc
    if path.suffix.lower() not in {".csv", ".json"}:
        raise ValueError(f"Unexpected table-source extension: {name}")
    return path


def load_sources(source: Path) -> list[tuple[dict, list[list[str]]]]:
    manifest = json.loads((source / "manifest.json").read_text(encoding="utf-8"))
    if manifest.get("schema_version") != 1:
        raise ValueError("Unsupported public-table manifest version")
    if manifest.get("table_count") != len(manifest["tables"]):
        raise ValueError("Manifest table count mismatch")
    seen, results = set(), []
    for entry in manifest["tables"]:
        table_id = entry["table_id"]
        if not re.fullmatch(r"Table_(?:1|S\d{2}[ABC]?)", table_id) or table_id in seen:
            raise ValueError(f"Invalid or duplicate table ID: {table_id}")
        seen.add(table_id)
        csv_path = safe_child(source, entry["csv_file"])
        metadata_path = safe_child(source, entry["metadata_file"])
        if sha256(csv_path) != entry["csv_sha256"] or sha256(metadata_path) != entry["metadata_sha256"]:
            raise ValueError(f"Public source checksum mismatch: {table_id}")
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        with csv_path.open(encoding="utf-8-sig", newline="") as stream:
            all_rows = list(csv.reader(stream))
        expected_columns = [f"column_{i + 1:02d}" for i in range(metadata["column_count"])]
        if not all_rows or all_rows[0] != expected_columns:
            raise ValueError(f"Storage header mismatch: {table_id}")
        rows = all_rows[1:]
        if len(rows) != metadata["row_count_including_display_headers"]:
            raise ValueError(f"Row count mismatch: {table_id}")
        if len(rows) != len(metadata["cell_layout"]) or any(len(r) != metadata["column_count"] for r in rows):
            raise ValueError(f"Table grid mismatch: {table_id}")
        for ri, cells in enumerate(metadata["cell_layout"]):
            occupied = set()
            for cell in cells:
                col, span = cell["column"], cell["colspan"]
                covered = set(range(col, col + span))
                if col < 0 or span < 1 or col + span > metadata["column_count"] or occupied & covered:
                    raise ValueError(f"Invalid cell span: {table_id} row {ri}")
                occupied |= covered
                if any(rows[ri][j] for j in range(col + 1, col + span)):
                    raise ValueError(f"Hidden text in merged-cell padding: {table_id} row {ri}")
        results.append((metadata, rows))
    return results


def notes_html(notes: list[str], class_name: str = "notes") -> str:
    return f'<div class="{class_name}">' + "".join(f"<p>{html.escape(note)}</p>" for note in notes) + "</div>"


def display_cells(metadata: dict, rows: list[list[str]]) -> list[tuple[int, int, str]]:
    return [(ri, cell["column"], rows[ri][cell["column"]])
            for ri, row in enumerate(metadata["cell_layout"]) for cell in row]


def section_html(metadata: dict, rows: list[list[str]], source: Path, output: Path) -> str:
    table_id = metadata["table_id"]
    label = metadata["title"]
    panel = f'<h3>{html.escape(metadata["panel_title"])}</h3>' if metadata["panel_title"] else ""
    parts = [f'<section class="table-section" id="{table_id}">', panel,
             f'<div class="downloads"><a href="{table_id}.csv">CSV</a>'
             f'<a href="{table_id}.json">Title, notes and provenance</a></div>',
             notes_html(metadata["notes_before"]), '<div class="table-wrap">',
             f'<table data-table-id="{table_id}"><caption>{html.escape(label)}</caption><thead>']
    header_rows = metadata["header_row_count"]
    for ri, row in enumerate(metadata["cell_layout"]):
        if ri == header_rows:
            parts.append("</thead><tbody>")
        group = len(row) == 1 and row[0]["colspan"] == metadata["column_count"]
        parts.append('<tr class="group-row">' if group else "<tr>")
        for cell in row:
            ci = cell["column"]
            value = rows[ri][ci]
            # Current archive has no vertical merges. Fail instead of silently
            # hiding values if a future manuscript introduces them.
            if cell.get("vertical_merge") is not None:
                raise ValueError(f"Vertical merge needs explicit rendering: {table_id}")
            is_header = ri < header_rows or ci == 0
            tag = "th" if is_header else "td"
            scope = ' scope="col"' if ri < header_rows else (' scope="row"' if ci == 0 else "")
            span = f' colspan="{cell["colspan"]}"' if cell["colspan"] > 1 else ""
            numeric = bool(re.match(r"^[\s+−\-<≤>≥(\[\d]", value)) and bool(re.search(r"\d", value))
            css = ' class="numeric"' if numeric and not is_header else ""
            parts.append(f'<{tag}{scope}{span}{css} data-row="{ri}" data-column="{ci}">{html.escape(value)}</{tag}>')
        parts.append("</tr>")
    parts.append("</tbody></table></div>")
    parts.append(notes_html(metadata["notes_after"] + metadata.get("shared_notes", [])))
    linked = metadata["linked_aggregate_results"]
    chain = "; ".join(Path(s["public_path"]).name for s in linked) or "No separate aggregate CSV is linked; values are documented in the archived manuscript."
    parts.append('<div class="provenance"><strong>Provenance: </strong>'
                 + html.escape(metadata["provenance_status"]) + ". "
                 + html.escape(chain) + "<br>" + html.escape(metadata["inference_scope"]) + "</div></section>")
    return "\n".join(parts)


def html_document(title: str, content: str) -> str:
    return '<!doctype html>\n<html lang="en"><head><meta charset="utf-8">' \
        '<meta name="viewport" content="width=device-width,initial-scale=1">' \
        f'<title>{html.escape(title)}</title><style>{CSS}</style></head><body>{content}</body></html>\n'


class CellReader(HTMLParser):
    """Read generated table cells back without comparing serialized HTML."""
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.table_id = None
        self.current_key = None
        self.fragments = []
        self.cells = {}

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attributes = dict(attrs)
        if tag == "table":
            self.table_id = attributes.get("data-table-id")
        if tag in {"th", "td"} and "data-row" in attributes:
            self.current_key = (self.table_id, int(attributes["data-row"]), int(attributes["data-column"]))
            self.fragments = []

    def handle_data(self, data: str) -> None:
        if self.current_key is not None:
            self.fragments.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag in {"th", "td"} and self.current_key is not None:
            if self.current_key in self.cells:
                raise ValueError("Duplicate rendered cell")
            self.cells[self.current_key] = "".join(self.fragments)
            self.current_key = None
        if tag == "table":
            self.table_id = None


def write_xlsx(tables: list[tuple[dict, list[list[str]]]], path: Path) -> dict:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.page import PageMargins
    except ImportError as exc:
        raise RuntimeError("--xlsx requires openpyxl; omit --xlsx for standard-library-only HTML and CSV") from exc
    book = Workbook()
    book.remove(book.active)
    for metadata, rows in tables:
        sheet = book.create_sheet(metadata["table_id"])
        width = metadata["column_count"]
        heading = metadata["title"] + ("\n" + metadata["panel_title"] if metadata["panel_title"] else "")
        sheet.append([heading])
        sheet.merge_cells(start_row=1, end_row=1, start_column=1, end_column=width)
        sheet["A1"].font = Font(name="Arial", size=12, bold=True)
        sheet["A1"].alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[1].height = 46
        sheet.append([SCOPE])
        sheet.merge_cells(start_row=2, end_row=2, start_column=1, end_column=width)
        sheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[2].height = 42
        for note in metadata["notes_before"]:
            sheet.append([note])
            n = sheet.max_row
            sheet.merge_cells(start_row=n, end_row=n, start_column=1, end_column=width)
            sheet.cell(n, 1).alignment = Alignment(wrap_text=True, vertical="top")
            sheet.row_dimensions[n].height = max(30, (len(note) // 150 + 1) * 15)
        start = sheet.max_row + 2
        for ri, row in enumerate(rows):
            for ci, value in enumerate(row):
                cell = sheet.cell(start + ri, ci + 1)
                cell.value = value
                cell.data_type = "s"  # Preserve numeric text and prevent formula interpretation.
                cell.font = Font(name="Arial", size=10, bold=ri < metadata["header_row_count"])
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if ri < metadata["header_row_count"]:
                    cell.fill = PatternFill("solid", fgColor="EEF3F6")
            for cell in metadata["cell_layout"][ri]:
                if cell["colspan"] > 1:
                    sheet.merge_cells(start_row=start + ri, end_row=start + ri,
                                      start_column=cell["column"] + 1, end_column=cell["column"] + cell["colspan"])
            sheet.row_dimensions[start + ri].height = max(27, (max(len(v) for v in row) // 34 + 1) * 14)
        for note in metadata["notes_after"] + metadata.get("shared_notes", []):
            sheet.append([note])
            n = sheet.max_row
            sheet.merge_cells(start_row=n, end_row=n, start_column=1, end_column=width)
            sheet.cell(n, 1).alignment = Alignment(wrap_text=True, vertical="top")
            sheet.row_dimensions[n].height = max(30, (len(note) // 150 + 1) * 15)
        from openpyxl.utils import get_column_letter
        for ci in range(1, width + 1):
            sheet.column_dimensions[get_column_letter(ci)].width = 36 if ci == 1 else 25
        sheet.freeze_panes = f"B{start + metadata['header_row_count']}"
        sheet.print_title_rows = f"{start}:{start + metadata['header_row_count'] - 1}"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins = PageMargins(left=.3, right=.3, top=.4, bottom=.4, header=.1, footer=.1)
    book.save(path)
    return {"path": path.name, "sha256": sha256(path), "worksheets": len(tables)}


def render(source: Path, output: Path, xlsx: bool = False) -> dict:
    if source.resolve() == output.resolve() or source.resolve() in output.resolve().parents:
        raise ValueError("Use an output directory outside the source-table directory")
    tables = load_sources(source)
    output.mkdir(parents=True, exist_ok=True)
    sections, checks, index_items = [], [], []
    expected_all = {}
    for metadata, rows in tables:
        table_id = metadata["table_id"]
        section = section_html(metadata, rows, source, output)
        page = html_document(metadata["title"],
            '<p class="navigation"><a href="index.html">All tables</a></p>'
            f'<p class="scope">{SCOPE}</p>' + section)
        page_path = output / (table_id + ".html")
        page_path.write_text(page, encoding="utf-8")
        for ext in ("csv", "json"):
            shutil.copy2(source / (table_id + "." + ext), output / (table_id + "." + ext))
        parser = CellReader()
        parser.feed(page)
        expected = {(table_id, ri, ci): value for ri, ci, value in display_cells(metadata, rows)}
        if parser.cells != expected:
            raise AssertionError(f"Rendered HTML differs from the public CSV: {table_id}")
        expected_all.update(expected)
        checks.append({"table_id": table_id, "physical_cells_checked": len(expected),
                       "html_csv_cell_match": True, "csv_sha256": sha256(output / (table_id + ".csv"))})
        sections.append(section)
        label = metadata["title"] + (" — " + metadata["panel_title"] if metadata["panel_title"] else "")
        index_items.append(f'<li data-search="{html.escape(label.lower(), quote=True)}">'
                           f'<a href="{table_id}.html">{html.escape(label)}</a> '
                           f'<span class="muted">[<a href="{table_id}.csv">CSV</a>]</span></li>')
    combined = html_document("All manuscript tables", f'<h1>All manuscript tables</h1><p class="scope">{SCOPE}</p>' + "\n".join(sections))
    combined_parser = CellReader()
    combined_parser.feed(combined)
    if combined_parser.cells != expected_all:
        raise AssertionError("Combined HTML differs from the public CSVs")
    (output / "all_tables.html").write_text(combined, encoding="utf-8")
    index = '<h1>Manuscript tables</h1>' + f'<p class="scope">{SCOPE}</p>' \
        + f'<p>{len(tables)} table panels: one main table and 34 supplementary panels.</p>' \
        + '<p><a href="all_tables.html">Open all tables for reading or printing</a></p>' \
        + ('<p><a href="manuscript_tables.xlsx">Download Excel workbook</a></p>' if xlsx else '') \
        + '<p class="search"><label for="filter">Find a table </label><input id="filter" type="search" placeholder="Table number, model, outcome..."></p>' \
        + '<ul id="table-list">' + "\n".join(index_items) + '</ul>' \
        + '<script>document.getElementById("filter").addEventListener("input",function(){const q=this.value.toLowerCase();document.querySelectorAll("#table-list li").forEach(x=>x.hidden=!x.dataset.search.includes(q));});</script>'
    (output / "index.html").write_text(html_document("Manuscript table index", index), encoding="utf-8")
    report = {"status": "PASS", "tables_rendered": len(tables), "source": "03_source_data/tables",
              "model_refitting": False, "participant_data_required": False, "docx_required": False,
              "physical_cells_verified": len(expected_all), "checks": checks,
              "outputs": ["index.html", "all_tables.html", "35 individual HTML tables", "35 CSV copies", "35 JSON metadata copies"]}
    if xlsx:
        report["workbook"] = write_xlsx(tables, output / "manuscript_tables.xlsx")
    (output / "render_verification.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--source", type=Path, help="Default: ROOT/03_source_data/tables")
    parser.add_argument("--output", "--output-dir", dest="output", type=Path, help="Default: ROOT/04_code/output/public_tables")
    parser.add_argument("--xlsx", action="store_true", help="Also build an Excel workbook; requires openpyxl")
    args = parser.parse_args()
    root = args.root.resolve()
    def location(value: Path | None, default: str) -> Path:
        return (value if value is not None and value.is_absolute() else root / (value or default)).resolve()
    report = render(location(args.source, "03_source_data/tables"), location(args.output, "04_code/output/public_tables"), args.xlsx)
    print(json.dumps({k: v for k, v in report.items() if k != "checks"}, indent=2))


if __name__ == "__main__":
    main()
